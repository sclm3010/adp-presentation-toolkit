from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "outputs" / "ADP论文汇报讲稿.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9E2EC"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)
        borders.append(tag)
    tc_pr.append(borders)


def add_run(paragraph, text, bold=False, italic=False, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_formula(doc, formula):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.2)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7FAFD")
    set_cell_border(cell, "B8CCE4")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(formula)
    r.font.name = "Cambria Math"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    r.font.size = Pt(14)


def add_section(doc, title, paragraphs):
    doc.add_heading(title, level=1)
    for item in paragraphs:
        if isinstance(item, tuple) and item[0] == "formula":
            add_formula(doc, item[1])
            continue
        if isinstance(item, tuple) and item[0] == "emphasis":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            add_run(p, item[1], bold=True, color="1F4D78")
            continue
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.25)
        p.add_run(item)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.line_spacing = 1.1
styles["Normal"].paragraph_format.space_after = Pt(6)

for name, size, color in [
    ("Heading 1", 16, "2E74B5"),
    ("Heading 2", 13, "2E74B5"),
    ("Heading 3", 12, "1F4D78"),
]:
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(4)
r = title.add_run("Value Iteration ADP 论文汇报讲稿")
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = RGBColor.from_string("0B2545")
r.font.name = "Calibri"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(18)
r = subtitle.add_run("Value Iteration Adaptive Dynamic Programming for Optimal Control of Discrete-Time Nonlinear Systems")
r.italic = True
r.font.size = Pt(10.5)
r.font.color.rgb = RGBColor.from_string("555555")

lead = doc.add_paragraph()
lead.paragraph_format.space_before = Pt(4)
lead.paragraph_format.space_after = Pt(12)
add_run(
    lead,
    "使用方式：这份讲稿按“向论文作者本人汇报阅读理解”的口吻整理。汇报时重点讲自己的理解，不评价论文贡献本身；如有不确定处，用“我目前理解是……”承接。",
    bold=True,
    color="1F4D78",
)

sections = [
    (
        "1. 开场",
        [
            "魏老师好，我这次主要汇报一下我对这篇 Value Iteration ADP 论文的阅读理解。因为这篇文章是您参与完成的工作，所以我会重点讲我理解到的整体思路、算法逻辑和几个关键结论。如果有理解不到位的地方，也想请您帮我指正。",
            "我目前的整体理解是：这篇论文研究的是离散时间非线性系统的无限时域最优控制问题。文中的 value iteration ADP 算法，重点不只是证明值函数最终能收敛，而是进一步讨论算法在有限次迭代后停止时，得到的控制律是否稳定、是否真的能用。",
        ],
    ),
    (
        "2. 问题建模",
        [
            "首先是问题建模。论文考虑的系统形式是：",
            ("formula", "x_{k+1}=F(x_k,u_k)"),
            "其中 x_k 是系统状态，u_k 是控制输入。控制目标是最小化无限时域性能指标：",
            ("formula", "J(x_0,u_0)=Σ_{k=0}^{∞} U(x_k,u_k)"),
            "也就是说，我们希望找到一个反馈控制律 u(x_k)，让系统稳定，同时让长期累积代价最小。",
            "理论上，这类问题可以通过动态规划或者 HJB 方程求解。但对于非线性系统来说，HJB 方程通常很难解析求解，而且动态规划会遇到维数灾难。所以文中采用 adaptive dynamic programming，也就是 ADP 的方法来近似求解。",
        ],
    ),
    (
        "3. 研究动机",
        [
            "然后是论文的研究动机。我理解这篇文章的出发点，是传统 value iteration 在实际应用中有几个限制。",
            "第一，很多传统 value iteration 算法从零初始值函数开始，也就是 V_0(x)=0。这个初始条件比较特殊，不够一般。",
            "第二，理论上最优控制律通常需要迭代到无穷次，也就是 i→∞。但在实际系统中，算法必须在有限次迭代后停止。",
            "第三，也是这篇论文最重要的地方：即使值函数看起来已经收敛，有限次迭代得到的控制律也不一定稳定，也不一定是 admissible 的。",
            ("emphasis", "所以我理解这篇论文真正要解决的问题是：value iteration ADP 不仅要能收敛，还要保证在有限步停止时，当前控制律是可靠的。"),
        ],
    ),
    (
        "4. 算法思路",
        [
            "接下来是算法思路。我理解文中的算法是从一个正半定的初始值函数开始：",
            ("formula", "V_0(x_k)=Φ(x_k)"),
            "然后每一轮迭代分两步。",
            "第一步，用当前值函数 V_i 求当前控制律：",
            ("formula", "v_i(x_k)=arg min_{u_k}{ U(x_k,u_k)+V_i(F(x_k,u_k)) }"),
            "这一步的含义是：当前控制输入不仅要考虑当前代价 U(x_k,u_k)，还要考虑下一状态之后的未来代价。",
            "第二步，用得到的控制律更新值函数：",
            ("formula", "V_{i+1}(x_k)=U(x_k,v_i(x_k))+V_i(F(x_k,v_i(x_k)))"),
            ("emphasis", "所以整个算法可以理解成：先根据当前值函数改进控制律，再根据控制律更新值函数。"),
            "这样不断迭代，最终希望值函数收敛到最优性能指标函数 J*(x)，控制律收敛到最优控制律。",
        ],
    ),
    (
        "5. 理论结果",
        [
            "然后是理论结果。我理解这里比较重要的结论是：初始值函数不一定必须取零。只要初始值函数是正半定的，迭代值函数最终都可以收敛到最优值函数：",
            ("formula", "lim_{i→∞} V_i(x_k)=J*(x_k)"),
            "不同初始值函数会影响收敛过程。比如有的情况下值函数是单调递增的，有的情况下是单调递减的，也有可能是非单调的。但最终都收敛到最优值函数。",
            "这一点相当于把传统 value iteration 的初始条件放宽了。",
        ],
    ),
    (
        "6. 终止准则与 admissibility",
        [
            "接下来是我认为这篇论文最关键的部分：终止准则和 admissibility。",
            "传统 value iteration 通常用收敛准则来停止，比如：",
            ("formula", "|V_{i+1}(x_k)-V_i(x_k)|≤ε"),
            "这个条件说明两次值函数之间的变化已经很小。但我理解文中强调的是，这个条件只能说明值函数数值上变化小，不能保证当前控制律 v_i(x) 一定稳定。",
            "因此，本文额外提出了 admissibility 终止条件：",
            ("formula", "V_{i+1}(x_k)-V_i(x_k)<U(x_k,v_i(x_k))"),
            "如果这个条件满足，文中通过后面的定理说明，当前迭代控制律 v_i(x) 是 admissible 的。也就是说，它可以稳定系统，并且对应的性能指标是有限的。",
            ("emphasis", "所以本文的终止逻辑不是只看“值函数是否收敛”，而是要同时考虑“控制律是否可用”。"),
        ],
    ),
    (
        "7. Algorithm 1 的逻辑",
        [
            "Algorithm 1 也围绕这个思想展开。算法先判断 V_1(x)≤V_0(x) 是否成立。",
            "如果成立，说明值函数是单调不增的，这种情况下算法进入 Block 1，主要检查收敛准则。",
            "如果不成立，算法进入 Block 2。这时只满足收敛准则还不够，还必须满足 admissibility 准则，才能停止并输出控制律。",
            ("emphasis", "我理解这个设计的意义是：它把 value iteration 从理论上的无限迭代，推进到了实际应用中的有限步可停止。"),
        ],
    ),
    (
        "8. 神经网络实现",
        [
            "然后是神经网络实现部分。论文中使用 neural networks 来实现 ADP 算法，主要有两个网络。",
            "一个是 critic network，用来近似值函数 V_i(x)。另一个是 action network，用来近似控制律 v_i(x)。",
            "critic 网络负责评价当前状态之后的长期代价，action 网络负责根据 critic 的评价选择控制输入。",
            "这里我理解为 ADP 中比较典型的“评价 + 改进”结构。",
        ],
    ),
    (
        "9. 仿真部分",
        [
            "仿真部分主要有两个例子。",
            "第一个是离散化倒立摆系统。文中选择了不同的正半定初始值函数，结果表明，值函数都可以收敛到最优值函数。而且在满足收敛准则和 admissibility 准则后，得到的控制律可以稳定系统。",
            "第二个是离散化扭转摆系统。这个例子更强调 admissibility 准则的必要性。文中展示了一个现象：如果只看传统收敛准则，算法可能已经满足停止条件，但此时得到的控制律并不稳定。也就是说，系统状态可能仍然发散。",
            "因此，仿真实验主要说明两点：第一，不同初始值函数下，本文算法仍然可以收敛；第二，只看值函数收敛不够，admissibility 终止准则是必要的。",
        ],
    ),
    (
        "10. 我的理解和后续计划",
        [
            "最后是我的理解。我目前理解这篇论文的核心点不是简单给出一个新的 value iteration 更新公式，而是解决传统 value iteration ADP 在实际应用中的一个关键问题：有限迭代停止后，当前控制律到底能不能用。",
            "传统方法更关注 i→∞ 时的最优性，但实际控制系统不可能等待无限次迭代。因此，我理解文中的 admissibility 分析和新的终止准则更贴近实际应用。",
            "如果后续要继续深入，我觉得可以从第一个倒立摆仿真实验开始复现。重点不是单纯画出值函数收敛曲线，而是对比两种情况：一种是只使用传统收敛准则；另一种是同时使用收敛准则和 admissibility 准则。然后观察两种情况下系统状态轨迹是否稳定。",
            "我的汇报到这里。以上是我目前对这篇论文主线的理解。如果这里面有理解不准确的地方，请老师帮我指正。后续如果需要，我也可以继续补充具体定理证明或者仿真实现。",
        ],
    ),
]

for title_text, body in sections:
    add_section(doc, title_text, body)

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("ADP 论文汇报讲稿")
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor.from_string("777777")

doc.save(OUT)
print(OUT)
